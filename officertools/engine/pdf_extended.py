"""Document operations that preserve untouched PDF pages and their text."""
import io
import math
import tempfile
from pathlib import Path

import fitz
from PIL import Image, ImageChops, ImageOps, ImageStat, ImageSequence
from pypdf import PdfReader, PdfWriter, Transformation

from .common import Result, safe_stem
from .pdf import _reader, _write, parse_ranges, _cjk_font


def open_pdf(job):
    doc = fitz.open(str(job.file()))
    if doc.needs_pass and not doc.authenticate(job.opt("password")):
        doc.close()
        raise RuntimeError("请输入正确的 PDF 打开密码。")
    return doc


def page_image(page, dpi):
    pix = page.get_pixmap(dpi=dpi, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def otsu(image):
    histogram = image.convert("L").histogram()
    count = sum(histogram)
    total = sum(i * n for i, n in enumerate(histogram))
    weight, subtotal, best, threshold = 0, 0, -1, 127
    for i, n in enumerate(histogram):
        weight += n
        subtotal += i * n
        if not weight or weight == count:
            continue
        variance = (subtotal / weight - (total - subtotal) / (count - weight)) ** 2 * weight * (count - weight)
        if variance > best:
            best, threshold = variance, i
    return threshold


def estimate_angle(image, maximum=10, check=lambda: None):
    maximum = max(0.2, min(45, float(maximum)))
    small = image.convert("L")
    small.thumbnail((900, 900))
    threshold = otsu(small)
    bw = small.point(lambda p: 0 if p <= threshold else 255)
    if ImageStat.Stat(bw).stddev[0] < 2:
        return 0.0
    def score(angle):
        check()
        turned = bw.rotate(angle, resample=Image.Resampling.NEAREST, fillcolor=255)
        rows = list(turned.resize((1, turned.height), Image.Resampling.BOX).getdata())
        return sum((255 - value) ** 2 for value in rows)
    coarse = [round(-maximum + i, 2) for i in range(int(maximum * 2) + 1)] + [0.0]
    angle = max(coarse, key=score)
    fine = [round(angle + i * 0.1, 2) for i in range(-6, 7) if abs(angle + i * .1) <= maximum]
    result = max(fine, key=score)
    return result if abs(result) >= .15 else 0.0


def _trim(image, margin):
    gray = image.convert("L")
    box = gray.point(lambda p: 255 if p < 245 else 0).getbbox()
    if not box:
        return image
    x0, y0, x1, y1 = box
    return image.crop((max(0, x0-margin), max(0, y0-margin), min(image.width, x1+margin), min(image.height, y1+margin)))


def scan_process(job):
    path = job.file()
    dpi = max(72, min(600, int(job.number("dpi", 200))))
    is_tiff = path.suffix.lower() in (".tif", ".tiff")
    source = Image.open(path) if is_tiff else open_pdf(job)
    count = getattr(source, "n_frames", 1) if is_tiff else len(source)
    chosen = set(parse_ranges(job.opt("ranges"), count))
    writer, frames, log = PdfWriter(), [], []
    reader = None if is_tiff else _reader(path, job.opt("password"))
    try:
        for n in range(count):
            job.checkpoint(n * 100 / count, "第 {} / {} 页".format(n + 1, count))
            if is_tiff:
                source.seek(n)
                image = source.convert("RGB")
            else:
                image = page_image(source[n], dpi) if n in chosen else None
            if n not in chosen:
                if is_tiff:
                    frames.append(image)
                else:
                    writer.add_page(reader.pages[n])
                continue
            if min(image.size) < job.number("min-edge", 0):
                if is_tiff:
                    frames.append(image)
                else:
                    writer.add_page(reader.pages[n])
                continue
            angle = 0
            if job.tool == "pdf-deskew":
                angle = estimate_angle(image, job.number("max-angle", 10), job.checkpoint)
                log.append("第 {} 页：校正角度 {:.1f}°".format(n + 1, angle))
                if not is_tiff and job.opt("mode", "preserve") == "preserve" and not job.flag("trim"):
                    page = reader.pages[n]
                    page.transfer_rotation_to_content()
                    # Image y points down; PDF y points up. Rotate about the page centre.
                    w, h = float(page.mediabox.width), float(page.mediabox.height)
                    page.add_transformation(Transformation().translate(-w/2, -h/2).rotate(angle).translate(w/2, h/2))
                    writer.add_page(page)
                    continue
                image = image.rotate(angle, Image.Resampling.BICUBIC, expand=True, fillcolor="white")
            else:
                gray = image.convert("L")
                mode = job.opt("bw-mode", "auto")
                if mode == "gray":
                    image = gray
                elif mode == "dither":
                    image = gray.convert("1", dither=Image.Dither.FLOYDSTEINBERG)
                else:
                    threshold = otsu(gray) if mode == "auto" else int(job.number("threshold", 160))
                    if not 0 <= threshold <= 255:
                        raise RuntimeError("阈值必须在 0 到 255 之间。")
                    image = gray.point(lambda p: 255 if p > threshold else 0, mode="1")
            if job.flag("trim"):
                image = _trim(image, max(0, int(job.number("margin", 12))))
            if is_tiff:
                frames.append(image)
            else:
                blob = io.BytesIO()
                image.save(blob, format="PDF", resolution=dpi, quality=max(10, min(100, int(job.number("quality", 92)))))
                blob.seek(0)
                writer.add_page(PdfReader(blob).pages[0])
        dest = job.outdir / (safe_stem(path) + ("_deskew" if job.tool == "pdf-deskew" else "_bw") + (".tiff" if is_tiff else ".pdf"))
        if is_tiff:
            frames[0].save(dest, save_all=True, append_images=frames[1:], compression="tiff_lzw", dpi=(dpi,dpi))
        else:
            _write(writer, dest)
        return Result("已处理 {} 页，原文件保留。\n".format(len(chosen)) + "\n".join(log), [dest])
    finally:
        source.close()


def watermark(job):
    from PIL import ImageDraw
    from reportlab.pdfgen.canvas import Canvas
    from reportlab.lib.utils import ImageReader
    reader = _reader(job.file(), job.opt("password"))
    writer = PdfWriter()
    selected = set(parse_ranges(job.opt("ranges"), len(reader.pages)))
    for n, page in enumerate(reader.pages):
        job.checkpoint(n * 100 / len(reader.pages))
        parity = job.opt("parity", "all")
        if n not in selected or (parity == "odd" and n % 2) or (parity == "even" and not n % 2):
            writer.add_page(page)
            continue
        w, h = float(page.mediabox.width), float(page.mediabox.height)
        stamp = io.BytesIO()
        canvas = Canvas(stamp, pagesize=(w, h))
        opacity = max(0, min(1, job.number("opacity", .18)))
        canvas.setFillAlpha(opacity)
        canvas.setStrokeAlpha(opacity)
        if job.opt("image"):
            im = Image.open(job.file("image")).convert("RGBA")
            sw = max(1, job.number("width", 180))
            sh = sw * im.height / im.width
        else:
            text = job.opt("text", "内部资料")
            size = max(6, min(300, int(job.number("size", 48))))
            font = _cjk_font(size * 3)
            box = font.getbbox(text)
            im = Image.new("RGBA", (max(1, box[2]-box[0]+12), max(1,box[3]-box[1]+12)))
            ImageDraw.Draw(im).text((6-box[0],6-box[1]),text,font=font,fill=job.opt("color", "#505050"))
            sw, sh = im.width/3, im.height/3
        blob = io.BytesIO()
        im.save(blob, "PNG")
        position = job.opt("position", "center")
        x = 24 + sw/2 if "left" in position else w-24-sw/2 if "right" in position else w/2
        y = 24 + sh/2 if "bottom" in position else h-24-sh/2 if "top" in position else h/2
        points = [(x,y)]
        if job.flag("tile"):
            points = [(x,y) for x in range(0,int(w)+1,max(80,int(sw+60))) for y in range(0,int(h)+1,max(80,int(sh+80)))]
        for x,y in points:
            canvas.saveState()
            canvas.translate(x,y)
            canvas.rotate(job.number("angle", 0))
            canvas.drawImage(ImageReader(blob),-sw/2,-sh/2,width=sw,height=sh,mask="auto")
            canvas.restoreState()
        canvas.save()
        stamp.seek(0)
        page.merge_page(PdfReader(stamp).pages[0])
        writer.add_page(page)
    dest = job.outdir / (safe_stem(job.file()) + "_watermark.pdf")
    return Result("已添加水印。", [_write(writer,dest)])


def page_master(job):
    reader = _reader(job.file(), job.opt("password"))
    order = job.opt("order") or "1-{}".format(len(reader.pages))
    writer = PdfWriter()
    for token in order.replace("，", ",").split(","):
        if token.strip().lower() in ("blank", "空白", "0"):
            p = reader.pages[0]
            writer.add_blank_page(width=float(p.mediabox.width),height=float(p.mediabox.height))
        else:
            for i in parse_ranges(token, len(reader.pages)):
                writer.add_page(reader.pages[i])
    dest = job.outdir / (safe_stem(job.file())+"_reordered.pdf")
    return Result("已编排 {} 页。".format(len(writer.pages)), [_write(writer,dest)])


def page_cut(job):
    reader = _reader(job.file(), job.opt("password"))
    rows, cols = int(job.number("rows",2)), int(job.number("cols",1))
    if not (1 <= rows <= 20 and 1 <= cols <= 20):
        raise RuntimeError("行列数需在 1 到 20 之间。")
    writer = PdfWriter()
    chosen = set(parse_ranges(job.opt("ranges"),len(reader.pages)))
    import copy
    for n,page in enumerate(reader.pages):
        job.checkpoint(n*100/len(reader.pages))
        if n not in chosen:
            writer.add_page(page)
            continue
        page.transfer_rotation_to_content()
        b=page.mediabox
        for row in range(rows):
            for col in range(cols):
                new=copy.deepcopy(page)
                from pypdf.generic import RectangleObject
                box=RectangleObject((float(b.left)+b.width*col/cols, float(b.top)-b.height*(row+1)/rows,
                                     float(b.left)+b.width*(col+1)/cols,float(b.top)-b.height*row/rows))
                new.mediabox=box
                new.cropbox=box
                writer.add_page(new)
    return Result("已切成 {} 页。".format(len(writer.pages)),[_write(writer,job.outdir/(safe_stem(job.file())+"_cut.pdf"))])


def add_links(job):
    with open_pdf(job) as doc:
        values=[float(v.strip()) for v in job.opt("rect","36,36,250,90").split(",")]
        if len(values)!=4 or values[2]<=values[0] or values[3]<=values[1]:
            raise RuntimeError("链接区域格式：左,上,右,下，单位为点。")
        url=job.opt("url")
        if not url:
            raise RuntimeError("请输入链接地址。")
        for i in parse_ranges(job.opt("ranges"),len(doc)):
            doc[i].insert_link({"kind":fitz.LINK_URI,"from":fitz.Rect(values),"uri":url})
        dest=job.outdir/(safe_stem(job.file())+"_links.pdf")
        doc.save(str(dest),garbage=4,deflate=True)
    return Result("已添加可点击链接。",[dest])


def to_html(job):
    with open_pdf(job) as doc:
        chosen=parse_ranges(job.opt("ranges"),len(doc))
        parts=[]
        for n in chosen:
            job.checkpoint(n*100/len(doc))
            parts.append(doc[n].get_text("html").replace('id="page0"','id="page{}"'.format(n+1)))
        from html import escape
        dest=job.outdir/(safe_stem(job.file())+".html")
        dest.write_text("<!doctype html><meta charset='utf-8'><title>"+escape(job.file().name)+"</title><style>body{background:#eee}div[id^=page]{position:relative;background:white;margin:20px auto;box-shadow:0 2px 10px #bbb;overflow:hidden}div[id^=page] p{position:absolute;margin:0;white-space:pre}div[id^=page] img{position:absolute}</style>"+"\n".join(parts),encoding="utf-8")
    return Result("已导出含文字位置、字体和图片的 HTML。",[dest])


def to_docx(job):
    from docx import Document
    from docx.shared import Pt
    from docx.enum.section import WD_SECTION_START
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    document=Document()
    with open_pdf(job) as doc:
        indexes=parse_ranges(job.opt("ranges"),len(doc))
        for k,n in enumerate(indexes):
            job.checkpoint(k*100/len(indexes),"转换第 {} 页".format(n+1))
            page=doc[n]
            section=document.sections[0] if k==0 else document.add_section(WD_SECTION_START.NEW_PAGE)
            section.page_width=Pt(page.rect.width)
            section.page_height=Pt(page.rect.height)
            for attr in ("left_margin","right_margin","top_margin","bottom_margin"):
                setattr(section,attr,Pt(24))
            # Reading order with embedded images and editable text/font runs.
            for block in page.get_text("dict",sort=True)["blocks"]:
                if block["type"]==1:
                    try:
                        document.add_picture(io.BytesIO(block["image"]),width=Pt(min(block["width"],page.rect.width-48)))
                    except Exception:
                        image=Image.open(io.BytesIO(block["image"])).convert("RGB")
                        blob=io.BytesIO(); image.save(blob,"PNG"); blob.seek(0)
                        document.add_picture(blob,width=Pt(min(block["width"],page.rect.width-48)))
                    continue
                p=document.add_paragraph()
                p.paragraph_format.space_after=Pt(3)
                for line in block.get("lines",[]):
                    for span in line["spans"]:
                        run=p.add_run(span["text"])
                        run.font.size=Pt(span["size"])
                        run.bold=bool(span["flags"] & 16)
                        run.italic=bool(span["flags"] & 2)
                        run.font.name=span["font"]
                    p.add_run("\n")
        dest=job.outdir/(safe_stem(job.file())+".docx")
        document.save(str(dest))
    return Result("已导出可编辑文字、字号和图片。复杂表格、分栏与浮动对象仍可能需要手动调整；扫描页不做 OCR。",[dest])
